import io
import zipfile

import formaty


def test_na_casti_po_odstavcich_a_nadpisech():
    odstavec = 'Rmutování je enzymatický proces. ' * 20   # ≈ 660 znaků
    text = '\n\n'.join(['# Úvod', odstavec, odstavec, '# Rmutování', odstavec, odstavec, odstavec, odstavec, odstavec])
    casti = formaty.na_casti(text, cil=2000)
    assert len(casti) >= 3
    assert all(len(c) <= 2000 + len(odstavec) for c in casti)
    # Nadpis začíná novou část, když je ta předchozí aspoň třetinová.
    assert any(c.startswith('# Rmutování') for c in casti)
    assert formaty.na_casti('') == []


def test_dekoduj_text_utf8_i_cp1250():
    assert formaty.dekoduj_text('Žluťoučký kůň'.encode('utf-8')) == 'Žluťoučký kůň'
    assert formaty.dekoduj_text('Žluťoučký kůň'.encode('cp1250')) == 'Žluťoučký kůň'


def test_docx_nadpisy_odstavce_a_tabulka():
    import docx

    d = docx.Document()
    d.add_heading('Chmelovar', level=1)
    d.add_paragraph('Var trvá 90 minut.')
    d.add_heading('Odpar', level=2)
    tab = d.add_table(rows=2, cols=2)
    tab.cell(0, 0).text, tab.cell(0, 1).text = 'Čas', 'Odpar'
    tab.cell(1, 0).text, tab.cell(1, 1).text = '90 min', '8 %'
    proud = io.BytesIO()
    d.save(proud)
    text = formaty.docx_na_text(proud.getvalue())
    assert '# Chmelovar' in text
    assert '## Odpar' in text
    assert 'Var trvá 90 minut.' in text
    assert '| Čas | Odpar |' in text and '| 90 min | 8 % |' in text


def _epub(kapitoly: list[tuple[str, str]]) -> bytes:
    proud = io.BytesIO()
    with zipfile.ZipFile(proud, 'w') as z:
        z.writestr('mimetype', 'application/epub+zip')
        z.writestr('META-INF/container.xml',
                   '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        manifest = ''.join(f'<item id="k{i}" href="k{i}.xhtml" media-type="application/xhtml+xml"/>' for i in range(len(kapitoly)))
        # Pořadí čtení schválně obrácené proti manifestu — rozhoduje spine.
        spine = ''.join(f'<itemref idref="k{i}"/>' for i in reversed(range(len(kapitoly))))
        z.writestr('OEBPS/content.opf',
                   f'<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
                   f'<manifest>{manifest}</manifest><spine>{spine}</spine></package>')
        for i, (nadpis, telo) in enumerate(kapitoly):
            z.writestr(f'OEBPS/k{i}.xhtml',
                       f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>x</title><style>p{{}}</style></head>'
                       f'<body><h1>{nadpis}</h1><p>{telo}</p><table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>')
    return proud.getvalue()


def test_epub_v_poradi_cteni_s_nadpisy_a_tabulkou():
    text = formaty.epub_na_text(_epub([('Slad', 'Ječmen se máčí.'), ('Rmut', 'Teplota 62 °C.')]))
    assert text.index('# Rmut') < text.index('# Slad')
    assert 'Ječmen se máčí.' in text and 'Teplota 62 °C.' in text
    assert '| A | B |' in text and '| 1 | 2 |' in text
    assert 'p{}' not in text   # obsah <style> se nepřepisuje


def test_csv_a_xlsx_na_tabulky():
    csv_data = 'Slad;Extrakt\nPlzeňský;81,5\nMnichovský;79,0\n'.encode('utf-8')
    casti = formaty.tabulky_na_casti(csv_data, 'text/csv')
    assert len(casti) == 1 and '| Slad | Extrakt |' in casti[0] and '| Plzeňský | 81,5 |' in casti[0]

    import openpyxl

    sesit = openpyxl.Workbook()
    lst = sesit.active
    lst.title = 'Voda'
    lst.append(['Ion', 'mg/l'])
    for i in range(45):
        lst.append([f'ion{i}', i])
    proud = io.BytesIO()
    sesit.save(proud)
    casti = formaty.tabulky_na_casti(proud.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert len(casti) == 2
    assert casti[0].startswith('## Voda (řádky 2–41)') and casti[1].startswith('## Voda (řádky 42–46)')
    assert all('| Ion | mg/l |' in c for c in casti)
